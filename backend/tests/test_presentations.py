from io import BytesIO
from pathlib import Path
from pptx import Presentation as PptxPresentation
from pptx.util import Inches


def create_sample_pptx_bytes(title: str = "Test Lecture", slide_count: int = 2) -> bytes:
    """Helper to generate an in-memory valid PPTX file for testing."""
    prs = PptxPresentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: Title slide
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = title
    s1.placeholders[1].text = "Introduction and Overview\nPresenter Name"

    if slide_count > 1:
        # Slide 2: Bullet points & notes
        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        s2.shapes.title.text = "Key Topics"
        tf = s2.placeholders[1].text_frame
        tf.text = "Point 1: Fundamental Principles"
        p2 = tf.add_paragraph()
        p2.text = "Point 2: Applications & Practice"
        p2.level = 0
        p3 = tf.add_paragraph()
        p3.text = "Sub-point 2a: Real-world examples"
        p3.level = 1

        notes = s2.notes_slide.notes_text_frame
        notes.text = "Teacher notes: Focus heavily on practical applications."

    if slide_count > 2:
        # Slide 3: Table
        s3 = prs.slides.add_slide(prs.slide_layouts[5])
        s3.shapes.title.text = "Comparison Table"
        table_shape = s3.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(8), Inches(2))
        table = table_shape.table
        table.cell(0, 0).text = "Method A"
        table.cell(0, 1).text = "Method B"
        table.cell(1, 0).text = "Fast execution"
        table.cell(1, 1).text = "High accuracy"

    buf = BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()


def test_upload_pptx_presentation_success(client, teacher_user, student_user):
    # 1. Create a course
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Physics 101", "description": "Mechanics and Waves"},
    )
    assert c_res.status_code == 201
    course_id = c_res.json()["id"]

    # 2. Upload PPTX presentation
    pptx_bytes = create_sample_pptx_bytes("Newtonian Mechanics", slide_count=3)
    res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": teacher_user.id},
        files={"file": ("mechanics_lecture.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
        data={"title": "Lecture 1: Newtonian Mechanics", "description": "First week slides"},
    )
    assert res.status_code == 201
    data = res.json()
    assert data["title"] == "Lecture 1: Newtonian Mechanics"
    assert data["description"] == "First week slides"
    assert data["presentation_type"] == "local"
    assert data["slide_count"] == 3
    assert len(data["slides"]) == 3
    assert data["slides"][0]["title"] == "Newtonian Mechanics"
    assert data["slides"][1]["title"] == "Key Topics"
    assert data["slides"][1]["notes"] == "Teacher notes: Focus heavily on practical applications."
    assert len(data["slides"][2]["tables"]) == 1
    assert data["download_url"] is not None

    # 3. Check course detail includes presentation
    detail_res = client.get(f"/courses/{course_id}", headers={"X-User-Id": student_user.id})
    assert detail_res.status_code == 200
    assert detail_res.json()["presentation_count"] == 1
    assert len(detail_res.json()["presentations"]) == 1


def test_upload_presentation_auto_title_from_filename(client, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Chemistry 101"},
    )
    course_id = c_res.json()["id"]

    pptx_bytes = create_sample_pptx_bytes("Organic Chemistry", slide_count=1)
    res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": teacher_user.id},
        files={"file": ("organic_chemistry_intro.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert res.status_code == 201
    assert res.json()["title"] == "organic chemistry intro"


def test_upload_presentation_rejects_invalid_extension(client, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Biology 101"},
    )
    course_id = c_res.json()["id"]

    res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": teacher_user.id},
        files={"file": ("notes.txt", b"Hello world", "text/plain")},
    )
    assert res.status_code == 400
    assert "Unsupported file format" in res.json()["detail"]


def test_upload_presentation_rejects_empty_file(client, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "History 101"},
    )
    course_id = c_res.json()["id"]

    res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": teacher_user.id},
        files={"file": ("empty.pptx", b"", "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert res.status_code == 400
    assert "Uploaded file is empty" in res.json()["detail"]


def test_student_cannot_upload_presentation(client, student_user, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Literature 101"},
    )
    course_id = c_res.json()["id"]

    pptx_bytes = create_sample_pptx_bytes("Shakespeare", slide_count=1)
    res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": student_user.id},
        files={"file": ("presentation.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert res.status_code == 403


def test_add_embed_presentation(client, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Art 101"},
    )
    course_id = c_res.json()["id"]

    res = client.post(
        f"/courses/{course_id}/presentations",
        headers={"X-User-Id": teacher_user.id},
        json={
            "title": "Renaissance Masters",
            "description": "Google Slides deck",
            "embed_url": "https://docs.google.com/presentation/d/12345abcde/edit#slide=id.p",
        },
    )
    assert res.status_code == 201
    data = res.json()
    assert data["title"] == "Renaissance Masters"
    assert data["presentation_type"] == "embed"
    assert "https://docs.google.com/presentation/d/12345abcde/embed" in data["embed_url"]


def test_download_presentation(client, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Math 101"},
    )
    course_id = c_res.json()["id"]

    pptx_bytes = create_sample_pptx_bytes("Calculus I", slide_count=2)
    upload_res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": teacher_user.id},
        files={"file": ("calculus.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    pres_id = upload_res.json()["id"]

    dl_res = client.get(f"/presentations/{pres_id}/download")
    assert dl_res.status_code == 200
    assert len(dl_res.content) == len(pptx_bytes)


def test_remove_presentation_and_disk_cleanup(client, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Computer Science 101"},
    )
    course_id = c_res.json()["id"]

    pptx_bytes = create_sample_pptx_bytes("Algorithms", slide_count=1)
    upload_res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": teacher_user.id},
        files={"file": ("algorithms.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    pres_id = upload_res.json()["id"]
    file_path = upload_res.json()["file_name"]

    del_res = client.delete(
        f"/courses/{course_id}/presentations/{pres_id}",
        headers={"X-User-Id": teacher_user.id},
    )
    assert del_res.status_code == 204

    # Verify presentation is removed
    get_res = client.get(f"/presentations/{pres_id}/slides")
    assert get_res.status_code == 404


def test_update_slide_mappings(client, teacher_user):
    c_res = client.post(
        "/courses",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Biology 101"},
    )
    course_id = c_res.json()["id"]

    # Add a video to the course
    v_res = client.post(
        f"/courses/{course_id}/videos",
        headers={"X-User-Id": teacher_user.id},
        json={"title": "Cell Biology Lecture", "youtube_url": "https://www.youtube.com/watch?v=dQw4w9WgXcQ"},
    )
    video_id = v_res.json()["id"]

    # Upload 3-slide presentation
    pptx_bytes = create_sample_pptx_bytes("Cell Division", slide_count=3)
    upload_res = client.post(
        f"/courses/{course_id}/presentations/upload",
        headers={"X-User-Id": teacher_user.id},
        files={"file": ("cells.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    pres_id = upload_res.json()["id"]

    # Map Slide 1 (index 0) to video_id with start=30, end=90
    map_res = client.put(
        f"/courses/{course_id}/presentations/{pres_id}/slide-mappings",
        headers={"X-User-Id": teacher_user.id},
        json={
            "mappings": [
                {"slide_index": 0, "video_id": video_id, "start": 30.0, "end": 90.0}
            ]
        },
    )
    assert map_res.status_code == 200
    slides = map_res.json()["slides"]
    assert len(slides) == 3
    assert slides[0]["video_id"] == video_id
    assert slides[0]["video_start"] == 30.0
    assert slides[0]["video_end"] == 90.0
    assert slides[1]["video_id"] is None
